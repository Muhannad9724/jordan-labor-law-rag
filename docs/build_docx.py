"""Markdown -> .docx for this project's docs.

Not a general Markdown implementation -- it handles exactly the constructs used in
EXPLAINER.md and LESSONS.md, and asserts at the end that nothing was silently
dropped. Deliberately not a plain text dump: headings, real Word tables, shaded
code blocks, bordered blockquotes and inline bold/italic/code all survive.

Arabic matters here. Word picks fonts for Arabic via the *complex script* slot,
which python-docx does not expose, so set_font() writes w:cs directly -- otherwise
the Arabic examples render in a mono font that has no Arabic glyphs.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

PROJECT = next((d for d in [Path.cwd(), *Path.cwd().parents, Path(__file__).resolve().parent.parent]
                if (d / "EXPLAINER.md").exists()), None)
assert PROJECT, "run from inside the project tree"

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
CODE_INK = RGBColor(0xA3, 0x1D, 0x1D)
ARABIC = re.compile(r"[\u0600-\u06FF]")

# code | bold | italic | link -- code first so ** inside a code span is left alone
INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|\*[^*\n]+?\*|\[[^\]]+\]\([^)]+\))")


def set_font(run, name, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    # Arabic runs must not be forced into a mono face with no Arabic coverage.
    rfonts.set(qn("w:cs"), BODY_FONT if ARABIC.search(run.text or "") else name)


def shade(par_or_cell, fill):
    el = par_or_cell._p.get_or_add_pPr() if hasattr(par_or_cell, "_p") else par_or_cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def border(par, edge="bottom", color="CCCCCC", sz="6", space="1"):
    pPr = par._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), sz)
    e.set(qn("w:space"), space)
    e.set(qn("w:color"), color)
    pbdr.append(e)


def add_inline(par, text, base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """Render **bold**, *italic*, `code` and [links] as separate runs."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 1:
            r = par.add_run(tok[1:-1])
            set_font(r, MONO_FONT, base_size - 1, base_bold, base_italic, CODE_INK)
        # Recurse: bold/italic frequently wrap an inline code span (**`max_tokens`**).
        # Emitting the inner text as one plain run would leave the backticks visible.
        elif tok.startswith("**") and tok.endswith("**"):
            add_inline(par, tok[2:-2], base_size, True, base_italic, base_color)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            add_inline(par, tok[1:-1], base_size, base_bold, True, base_color)
        elif tok.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            r = par.add_run(m.group(1))
            set_font(r, BODY_FONT, base_size, base_bold, base_italic, ACCENT)
            r.underline = True
        else:
            r = par.add_run(tok)
            set_font(r, BODY_FONT, base_size, base_bold, base_italic, base_color)


def build(md_path: Path, out_path: Path, subtitle: str):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.15
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)

    counts = {"h": 0, "p": 0, "table": 0, "code": 0, "quote": 0, "li": 0, "hr": 0}
    i = 0
    first_h1_used = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            shade(p, "F4F5F7")
            r = p.add_run("\n".join(buf))
            set_font(r, MONO_FONT, 9)
            counts["code"] += 1
            continue

        # pipe table
        if stripped.startswith("|") and stripped.endswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in block]
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.paragraphs[0].text = ""
                        par = cell.paragraphs[0]
                        par.paragraph_format.space_after = Pt(2)
                        add_inline(par, row[ci] if ci < len(row) else "",
                                   base_size=9.5, base_bold=(ri == 0))
                        if ri == 0:
                            shade(cell, "E8EDF3")
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
                counts["table"] += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            border(p, "bottom", "B8C4D0", "8")
            counts["hr"] += 1
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1 and not first_h1_used:
                first_h1_used = True
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(re.sub(r"[*`]", "", text))
                set_font(r, BODY_FONT, 22, True, color=ACCENT)
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(14)
                border(sp, "bottom", "B8C4D0", "8")
                sr = sp.add_run(subtitle)
                set_font(sr, BODY_FONT, 10.5, italic=True, color=MUTED)
            else:
                h = doc.add_heading(level=min(level, 4))
                h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
                h.paragraph_format.space_after = Pt(4)
                for r in list(h.runs):
                    r.text = ""
                add_inline(h, text, base_size={1: 18, 2: 15, 3: 12.5, 4: 11}[level],
                           base_bold=True, base_color=ACCENT)
            counts["h"] += 1
            i += 1
            continue

        # blockquote (consecutive > lines).
        # Source is hard-wrapped at ~80 cols, so one logical quote paragraph spans
        # several lines. Render line-by-line and a **bold** span that straddles a
        # wrap point never matches -- join first, split on blank > lines.
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]).rstrip())
                i += 1
            groups, cur = [], []
            for q in buf:
                if q.strip():
                    cur.append(q.strip())
                elif cur:
                    groups.append(" ".join(cur))
                    cur = []
            if cur:
                groups.append(" ".join(cur))
            for j, g in enumerate(groups):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(3 if j < len(groups) - 1 else 10)
                border(p, "left", "1F4E79", "18", "8")
                shade(p, "F7F9FB")
                add_inline(p, g, base_size=10.5)
            counts["quote"] += 1
            continue

        # list item, plus any wrapped continuation lines (same reason as above)
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent, marker, text = m.group(1), m.group(2), m.group(3)
            i += 1
            parts = [text.strip()]
            while i < len(lines) and lines[i].strip() and not re.match(
                r"^\s*(#{1,4}\s|[-*]\s|\d+\.\s|>|\||```|-{3,}$)", lines[i]
            ):
                parts.append(lines[i].strip())
                i += 1
            ordered = bool(re.match(r"\d+\.", marker))
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * (len(indent) // 2))
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, " ".join(parts))
            counts["li"] += 1
            continue

        if not stripped:
            i += 1
            continue

        # paragraph: join until a blank line or a new block construct
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^\s*(#{1,4}\s|[-*]\s|\d+\.\s|>|\||```|-{3,}$)", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, " ".join(buf))
        counts["p"] += 1

    doc.save(out_path)
    return counts


JOBS = [
    ("EXPLAINER.md", "How this project works, explained from scratch",
     "Jordan Labour Law RAG - Explainer.docx"),
    ("LESSONS.md", "Mistakes made across 23 phases, and what each one taught",
     "Jordan Labour Law RAG - Lessons.docx"),
]

for src, subtitle, dest in JOBS:
    md = PROJECT / src
    out = PROJECT / "docs" / dest
    out.parent.mkdir(exist_ok=True)
    c = build(md, out, subtitle)
    kb = out.stat().st_size / 1024
    print(f"{dest}\n   {kb:6.0f} KB  " + "  ".join(f"{k}={v}" for k, v in c.items()))

    # nothing silently dropped, and no raw Markdown left showing
    raw = md.read_text(encoding="utf-8")
    assert c["code"] == raw.count("\n```") // 2, f"code blocks: {c['code']}"
    assert c["table"] > 0, "tables lost"

    chk = Document(out)
    txt = "\n".join([p.text for p in chk.paragraphs]
                    + [cl.text for t in chk.tables for r in t.rows for cl in r.cells])
    for marker in ("**", "`"):
        assert marker not in txt, f"unconverted {marker!r} in {dest}"
    if ARABIC.search(raw):
        assert ARABIC.search(txt), f"Arabic lost in {dest}"
    print(f"   verified: no raw Markdown left, {len(txt.split())} words")
print("\nok")
